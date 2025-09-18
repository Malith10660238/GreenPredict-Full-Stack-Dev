import os
from docx import Document

# Path to your flutter lib folder
root_dir = "lib"

# Output file names
txt_file = "all_code.txt"
docx_file = "all_code.docx"

# --- TXT Export ---
with open(txt_file, "w", encoding="utf-8") as outfile:
    for foldername, subfolders, filenames in os.walk(root_dir):
        for filename in sorted(filenames):
            if filename.endswith(".dart"):
                file_path = os.path.join(foldername, filename)
                outfile.write(f"\n\n===== File: {file_path} =====\n\n")
                with open(file_path, "r", encoding="utf-8") as infile:
                    outfile.write(infile.read())

print(f"✅ Exported TXT: {txt_file}")

# --- DOCX Export ---
doc = Document()
for foldername, subfolders, filenames in os.walk(root_dir):
    for filename in sorted(filenames):
        if filename.endswith(".dart"):
            file_path = os.path.join(foldername, filename)
            doc.add_heading(f"File: {file_path}", level=2)
            with open(file_path, "r", encoding="utf-8") as infile:
                doc.add_paragraph(infile.read())

doc.save(docx_file)
print(f"✅ Exported DOCX: {docx_file}")
